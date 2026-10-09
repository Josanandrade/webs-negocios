"""Utilidades para construir datos de prueba reales (PDFs reales, filas reales)."""
import hashlib
import uuid

import pymupdf as fitz
from sqlalchemy import text

from app.db import user_session


def make_pdf(pages: list[str]) -> bytes:
    doc = fitz.open()
    for body in pages:
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(56, 56, 540, 790), body, fontsize=11, fontname="helv")
    data = doc.tobytes()
    doc.close()
    return data


def insert_document_with_pages(user_id: uuid.UUID, page_texts: list[str]) -> tuple[uuid.UUID, list[uuid.UUID]]:
    """Inserta un documento ya "extraído" directamente en BD (para tests del modelo)."""
    doc_id = uuid.uuid4()
    with user_session(user_id) as s:
        s.execute(
            text("insert into documents (id, user_id, title, original_filename, mime_type, size_bytes, sha256, storage_key, status, page_count)"
                 " values (:id, :u, 'Temario', 'temario.pdf', 'application/pdf', 100, :sha, :key, 'ready', :n)"),
            {"id": doc_id, "u": user_id, "sha": hashlib.sha256(doc_id.bytes).hexdigest(),
             "key": f"{user_id}/{doc_id}.pdf", "n": len(page_texts)},
        )
        page_ids = []
        for i, body in enumerate(page_texts, start=1):
            pid = s.execute(
                text("insert into document_pages (document_id, user_id, page_number, text, char_count, extraction_method, quality_score, is_eligible)"
                     " values (:d, :u, :n, :t, :c, 'native', 1, true) returning id"),
                {"d": doc_id, "u": user_id, "n": i, "t": body, "c": len(body)},
            ).scalar_one()
            page_ids.append(pid)
    return doc_id, page_ids


def insert_question(session, user_id, document_id, page_id, page_number, *, options=None, sources=True,
                    stem="¿Cuál es la dosis diaria del medicamento A?"):
    """Inserta una pregunta completa. `options` = lista de (label, texto, es_correcta)."""
    options = options if options is not None else [
        ("A", "20 mg", True), ("B", "40 mg", False), ("C", "10 mg", False), ("D", "5 mg", False)]
    qid = session.execute(
        text("insert into questions (user_id, document_id, stem, question_type, difficulty, status, fingerprint)"
             " values (:u, :d, :stem, 'datum', 'medium', 'auto_validated', :fp) returning id"),
        {"u": user_id, "d": document_id, "stem": stem, "fp": hashlib.sha1(stem.encode()).hexdigest()},
    ).scalar_one()
    option_ids = {}
    for label, body, correct in options:
        option_ids[label] = session.execute(
            text("insert into question_options (question_id, user_id, label, text, is_correct)"
                 " values (:q, :u, :l, :t, :c) returning id"),
            {"q": qid, "u": user_id, "l": label, "t": body, "c": correct},
        ).scalar_one()
    if sources:
        session.execute(
            text("insert into question_sources (question_id, user_id, document_id, role, page_id, page_number, quote)"
                 " values (:q, :u, :d, 'answer', :p, :n, 'El medicamento A se administra a dosis de 20 mg')"),
            {"q": qid, "u": user_id, "d": document_id, "p": page_id, "n": page_number},
        )
    return qid, option_ids


def make_scanned_pdf(pages: list[str], dpi: int = 200) -> bytes:
    """PDF 'escaneado' real: cada página es solo una imagen del texto, sin capa de texto."""
    src = fitz.open(stream=make_pdf(pages), filetype="pdf")
    out = fitz.open()
    for page in src:
        pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, stream=pix.tobytes("png"))
    data = out.tobytes()
    out.close()
    src.close()
    return data


def make_noise_pdf() -> bytes:
    """Página con una imagen de ruido: el OCR no puede extraer texto fiable."""
    import random

    rng = random.Random(42)
    pix = fitz.Pixmap(fitz.csGRAY, fitz.IRect(0, 0, 600, 800), False)
    pix.set_rect(pix.irect, (255,))
    for _ in range(25000):
        x, y = rng.randrange(600), rng.randrange(800)
        pix.set_pixel(x, y, (0,))
    out = fitz.open()
    page = out.new_page()
    page.insert_image(page.rect, stream=pix.tobytes("png"))
    data = out.tobytes()
    out.close()
    return data


def merge_pdfs(*parts: bytes) -> bytes:
    out = fitz.open()
    for part in parts:
        with fitz.open(stream=part, filetype="pdf") as src:
            out.insert_pdf(src)
    data = out.tobytes()
    out.close()
    return data


def make_docx(paragraphs: list[str], page_break_after: set[int] = frozenset()) -> bytes:
    import io

    from docx import Document
    from docx.enum.text import WD_BREAK

    d = Document()
    for i, p in enumerate(paragraphs):
        para = d.add_paragraph(p)
        if i in page_break_after:
            para.add_run().add_break(WD_BREAK.PAGE)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_structured_pdf(pages: list[list[tuple[str, str]]], toc: list[list] | None = None) -> bytes:
    """PDF con tipografía real: elementos ('h1'|'h2'|'p', texto) por página. `toc` opcional
    en formato PyMuPDF [[nivel, título, página], ...]."""
    styles = {"h1": (18, "hebo"), "h2": (14, "hebo"), "p": (11, "helv")}
    doc = fitz.open()
    for elements in pages:
        page = doc.new_page()
        y = 56.0
        for kind, body in elements:
            size, font = styles[kind]
            rect = fitz.Rect(56, y, 540, 800)
            unused = page.insert_textbox(rect, body, fontsize=size, fontname=font)
            assert unused >= 0, "el texto no cabe en la página de prueba"
            y = 800 - unused + size * 0.8
    if toc:
        doc.set_toc(toc)
    data = doc.tobytes()
    doc.close()
    return data
